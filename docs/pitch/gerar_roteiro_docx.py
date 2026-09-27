#!/usr/bin/env python3
"""Gera docs/pitch/roteiro_pitch.docx a partir do roteiro do pitch de 5 minutos."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

PREDICTA = RGBColor(0x4F, 0x46, 0xE5)
LARANJA = RGBColor(0xDC, 0x6B, 0x18)
CINZA = RGBColor(0x64, 0x74, 0x8B)

SECOES = [
    (
        "Slide 1 · Problema", "0:00 – 0:40",
        [
            "Todo verão a história se repete: uma onda de calor, todo mundo liga o ar-condicionado na mesma hora, "
            "e o sistema elétrico inteiro é levado ao limite — não pelo consumo total do dia, mas por poucas horas de pico.",
            "Demanda não é consumo. Dois dias podem gastar a mesma energia, mas um deles concentra tudo entre 18h e 21h. "
            "É esse pico que dimensiona rede, geração e reserva — e que custa caro para a distribuidora.",
            "E o consumidor? Ele não tem nenhum incentivo para sair do pico, porque a tarifa que ele vê é plana. "
            "O sistema é dimensionado pelo pior dia, mas cobrado como se todo dia fosse igual.",
        ],
    ),
    (
        "Slide 2 · Solução: os dois motores", "0:40 – 1:50",
        [
            "A Predicta resolve isso com dois motores conectados por um único sinal.",
            "À esquerda, só dados oficiais: carga, geração e programação diária do ONS, tarifas da ANEEL, "
            "dez anos de clima em grade de dez quilômetros e a localização real de usinas e subestações.",
            "O Motor 1 aprende com tudo isso e prevê a demanda das próximas 24 horas, região por região, "
            "com faixa de confiança — e mede a pressão do sistema: quanta folga existe entre o que foi programado e o que vai acontecer.",
            "Esse resultado vira um sinal horário auditável, que alimenta o Motor 2: o motor de tarifa dinâmica. "
            "E aqui está o ponto que interessa à banca: o Motor 2 tem regras de proteção travadas. "
            "A tarifa nunca sai da faixa de 85% a 130% da tarifa base, nunca varia mais de 10% de uma hora para a outra, "
            "a fatura média do mês é neutra por construção — não é aumento disfarçado de receita — e, se qualquer dado falhar, a tarifa volta a ser a normal.",
            "Na ponta direita, dois beneficiários: o consumidor descobre quando a energia está mais barata, "
            "e a distribuidora vê sua curva achatar.",
        ],
    ),
    (
        "Slide 3 · Demo ao vivo", "1:50 – 2:30",
        [
            "Deixa eu mostrar isso funcionando. [ABRIR /produto/]",
            "Este é o mapa real das concessões da ANEEL. Escolho a [DISTRIBUIDORA] e simulo um dia real de operação. "
            "[SIMULAR — janela de replay ensaiada]",
            "Três curvas: o que o ONS programou, o que a Predicta previu e o que de fato aconteceu — a nossa previsão colada no realizado.",
            "Para este cliente residencial, deslocar uma parte flexível do consumo vale [R$ X POR MÊS] na conta.",
            "E na aba da carteira: com [Y]% de adesão em [N] clientes, o pico da carteira cai [Z] MW. "
            "Isso, escalado para a base inteira da distribuidora, é capacidade de rede liberada na hora mais crítica do dia.",
        ],
    ),
    (
        "Slide 4 · Confiança e rastreabilidade", "2:30 – 3:10",
        [
            "E dá para confiar nesses números? Três razões.",
            "Primeira: só usamos fontes oficiais — ONS, ANEEL e bases climáticas públicas. Nada foi inventado.",
            "Segunda: cada tarifa emitida carrega uma espécie de nota fiscal de dados — o carimbo de onde veio cada insumo. "
            "Qualquer valor que aparecer na tela pode ser reproduzido e auditado depois, item por item.",
            "Terceira: o modelo foi congelado e confrontado com 30 dias reais que ele nunca tinha visto. "
            "Resultado: erro de 2,3% contra 4,5% do método tradicional no Sudeste — metade do erro. "
            "No Sul, 3,3% contra 6,7%. E nas horas de calor extremo, quando mais importa, o erro se mantém em cerca de 2%.",
        ],
    ),
    (
        "Slide 5 · Modelo de negócio", "3:10 – 3:50",
        [
            "Quem paga? A distribuidora: uma taxa de instalação e um fee de serviço recorrente.",
            "O que ela ganha? A curva achatada — vende mais energia nos vales, sofre menos no pico, "
            "e adia investimento em rede que só existe para atender poucas horas por ano.",
            "O consumidor entra de graça: recebe a previsão de 24 horas pela própria distribuidora e economiza aderindo à tarifa dinâmica. "
            "Quem quiser mais — horizonte estendido, otimização automática — paga um fee premium opcional para nós.",
            "Todo mundo ganha: distribuidora, consumidor e o próprio sistema, que despacha menos térmica cara.",
        ],
    ),
    (
        "Slide 6 · Viabilidade e roadmap", "3:50 – 4:25",
        [
            "Onde estamos? Isso que vocês viram não é mockup: o pipeline roda de ponta a ponta hoje, "
            "com quatro subsistemas do SIN e três anos de dados reais, e uma API pronta para integrar ao billing da distribuidora.",
            "Sendo honestos sobre o que falta: colocar a previsão meteorológica ao vivo em operação — o código já existe — e a operação agendada 24/7.",
            "O caminho natural é um piloto com uma distribuidora em ambiente regulatório experimental: "
            "primeiro em modo sombra, com a tarifa simulada rodando em paralelo à real; depois opt-in com clientes voluntários; depois escala.",
        ],
    ),
    (
        "Slide 7 · Inovação", "4:25 – 4:50",
        [
            "Por que ninguém fez isso ainda? As ferramentas que existem hoje — bandeiras tarifárias, tarifa branca — são estáticas e reativas. "
            "A Predicta é preditiva, horária, localizada por região e com proteções para os dois lados da relação.",
            "E a arquitetura de dois motores desacoplados significa que a distribuidora pluga o motor de tarifa no sistema que já tem, "
            "sem reconstruir nada.",
        ],
    ),
    (
        "Slide 8 · Fechamento", "4:50 – 5:00",
        [
            "A demanda vai continuar mudando com o clima. A tarifa precisa aprender a acompanhar.",
            "Somos a Predicta — e buscamos uma distribuidora parceira para o piloto. Obrigada.",
        ],
    ),
]

OBSERVACOES = [
    "Placeholders [R$ X], [Y], [Z], [N] e [DISTRIBUIDORA]: preencher com os valores do cenário de demo congelado "
    "(ensaiar a simulação e anotar os números exibidos na tela).",
    "Números de validação citados (verificados no repositório): SE/CO WAPE 2,3% (E3/XGBoost) vs 4,5% (baseline), "
    "Sul 3,3% vs 6,7%, calor extremo ≈2,0% — holdout de 720h.",
    "Ritmo alvo: ~150 palavras/minuto. Ensaiar 3× com cronômetro; a demo de 40s é o maior risco — ter screenshot de backup nos slides.",
    "Se a demo falhar: seguir com o screenshot do slide 3 e citar os números memorizados sem quebrar a narrativa.",
]


def main() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run("Predicta · Roteiro do Pitch — 5 minutos")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = PREDICTA

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Demanda energética, clima e operação do sistema · Hackathon 2026")
    run.font.color.rgb = CINZA
    run.font.size = Pt(11)

    for nome, tempo, falas in SECOES:
        h = doc.add_paragraph()
        h.space_before = Pt(14)
        run = h.add_run(f"{nome}  ")
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = PREDICTA
        run = h.add_run(f"[{tempo}]")
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = LARANJA
        for fala in falas:
            p = doc.add_paragraph(fala)
            p.paragraph_format.space_after = Pt(6)

    h = doc.add_paragraph()
    run = h.add_run("Observações de preparação")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = PREDICTA
    for obs in OBSERVACOES:
        doc.add_paragraph(obs, style="List Bullet")

    out = Path(__file__).with_name("roteiro_pitch.docx")
    doc.save(out)
    print(f"OK: {out}")


if __name__ == "__main__":
    main()
